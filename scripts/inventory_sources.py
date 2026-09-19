#!/usr/bin/env python3
"""Inventory TRAIN Parquet metadata and locate proposal variable references."""

from __future__ import annotations

import argparse
import csv
import json
import platform
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, TextIO

import docx
import pyarrow
import pyarrow.parquet as pq


TASK_ID = "TASK_06"
SCRIPT_VERSION = "1.0.0"
EXPECTED_TRAIN_FILE_COUNT = 32
STRUCTURAL_NAMES = {
    "case_id",
    "target",
    "date_decision",
    "MONTH",
    "WEEK_NUM",
    "num_group1",
    "num_group2",
}
TECHNICAL_ROLES = {
    "case_id": "JOIN_KEY",
    "target": "LABEL",
    "date_decision": "TIME_METADATA",
    "MONTH": "TIME_METADATA",
    "WEEK_NUM": "TIME_METADATA",
    "num_group1": "GROUP_INDEX",
    "num_group2": "GROUP_INDEX",
}
RAW_FIELD_TOKEN_RE = re.compile(
    r"(?<![A-Za-z0-9_])([A-Za-z][A-Za-z0-9_]*_[0-9]+[A-Z]?)(?![A-Za-z0-9_])"
)
FILENAME_RE = re.compile(
    r"^train_(?P<source>.+?)_(?P<depth>[012])(?:_(?P<shard>[0-9]+))?\.parquet$"
)
REPORT_FILENAMES = (
    "table_inventory.csv",
    "field_inventory.csv",
    "proposal_candidates.csv",
    "inventory_summary.json",
)
TABLE_HEADERS = (
    "file_name",
    "relative_path",
    "file_size_bytes",
    "metadata_rows",
    "row_group_count",
    "column_count",
    "provisional_table_group",
    "inferred_depth",
    "inferred_shard",
    "naming_status",
    "columns_json",
    "arrow_types_json",
    "has_case_id",
    "case_id_arrow_type",
    "has_target",
    "has_date_decision",
    "has_MONTH",
    "has_WEEK_NUM",
    "has_num_group1",
    "has_num_group2",
    "group_file_count",
    "group_metadata_row_total",
    "group_ordered_schema_match",
)
FIELD_HEADERS = (
    "file_name",
    "provisional_table_group",
    "column_position_zero_based",
    "field_name",
    "arrow_type",
    "dictionary_status",
    "dictionary_description",
    "technical_role",
    "substantive_assignment_status",
)
CANDIDATE_HEADERS = (
    "candidate_name",
    "in_dictionary",
    "dictionary_description",
    "observed_in_train",
    "location_status",
    "physical_file_count",
    "provisional_group_count",
    "physical_files_json",
    "provisional_groups_json",
    "arrow_types_json",
    "proposal_reference_locations_json",
    "proposal_contexts_json",
    "is_structural_reference",
    "researcher_review_status",
)
UNPERFORMED_CHECKS = (
    "Row-value inspection",
    "case_id uniqueness within non-base tables",
    "Key overlap across shards",
    "Application coverage of source tables",
    "Missing-value rates",
    "Value distributions",
    "Applicant/person/contract role validation",
    "Point-in-time availability",
    "Monetary units and transformations",
    "Substantive T/AD assignment",
    "Final model-variable eligibility",
)


class DictionaryValidationError(ValueError):
    """Raised when dictionary variable names are blank or duplicated."""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--train-dir", required=True, type=Path)
    parser.add_argument("--dictionary", required=True, type=Path)
    parser.add_argument("--proposal", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser.parse_args()


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def csv_bool(value: bool) -> str:
    return "true" if value else "false"


def provisional_filename_group(file_name: str) -> dict[str, Any]:
    if file_name == "train_base.parquet":
        return {
            "provisional_table_group": "train_base",
            "inferred_depth": None,
            "inferred_shard": None,
            "naming_status": "BASE",
            "naming_warning": None,
        }
    match = FILENAME_RE.fullmatch(file_name)
    if match:
        depth = int(match.group("depth"))
        shard_text = match.group("shard")
        return {
            "provisional_table_group": f"train_{match.group('source')}_{depth}",
            "inferred_depth": depth,
            "inferred_shard": int(shard_text) if shard_text is not None else None,
            "naming_status": "EXPECTED_PATTERN",
            "naming_warning": None,
        }
    return {
        "provisional_table_group": Path(file_name).stem,
        "inferred_depth": None,
        "inferred_shard": None,
        "naming_status": "UNEXPECTED_PATTERN",
        "naming_warning": (
            f"{file_name} does not match train_<source>_<depth>[_<shard>].parquet "
            "or train_base.parquet; no depth or shard was inferred."
        ),
    }


def validate_dictionary_rows(
    rows: Iterable[dict[str, str]],
) -> tuple[list[dict[str, str]], dict[str, str], dict[str, Any]]:
    preserved: list[dict[str, str]] = []
    blank_rows: list[int] = []
    names: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        name = row.get("Variable")
        description = row.get("Description")
        name = "" if name is None else name
        description = "" if description is None else description
        preserved.append({"Variable": name, "Description": description})
        names.append(name)
        if not name.strip():
            blank_rows.append(row_number)

    name_counts = Counter(names)
    duplicate_names = sorted(
        name for name, count in name_counts.items() if name.strip() and count > 1
    )
    findings = {
        "blank_variable_name_row_numbers": blank_rows,
        "duplicate_variable_names": duplicate_names,
        "empty_description_variable_names": sorted(
            row["Variable"]
            for row in preserved
            if row["Variable"].strip() and row["Description"] == ""
        ),
    }
    if blank_rows or duplicate_names:
        raise DictionaryValidationError(
            "Dictionary variable names are ambiguous: "
            f"blank rows={blank_rows}, duplicate names={duplicate_names}"
        )
    mapping = {row["Variable"]: row["Description"] for row in preserved}
    return preserved, mapping, findings


def read_dictionary(path: Path) -> tuple[list[dict[str, str]], dict[str, str], dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames or []
        missing = [name for name in ("Variable", "Description") if name not in fieldnames]
        if missing:
            raise DictionaryValidationError(
                f"Dictionary is missing required columns: {missing}"
            )
        return validate_dictionary_rows(reader)


def exact_token_pattern(name: str) -> re.Pattern[str]:
    return re.compile(
        rf"(?<![A-Za-z0-9_]){re.escape(name)}(?![A-Za-z0-9_])"
    )


def short_context(text: str, name: str, max_chars: int = 240) -> str:
    match = exact_token_pattern(name).search(text)
    if not match:
        return ""
    padding = max_chars // 2
    start = max(0, match.start() - padding)
    end = min(len(text), match.end() + padding)
    excerpt = re.sub(r"\s+", " ", text[start:end]).strip()
    if start:
        excerpt = "…" + excerpt
    if end < len(text):
        excerpt = excerpt + "…"
    return excerpt[:max_chars]


def proposal_text_blocks(document: Any) -> tuple[list[dict[str, str]], dict[str, int]]:
    blocks: list[dict[str, str]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        blocks.append(
            {
                "location": f"paragraph:{paragraph_number}",
                "text": paragraph.text,
            }
        )

    table_cell_paragraph_count = 0
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            for cell_number, cell in enumerate(row.cells, start=1):
                for paragraph_number, paragraph in enumerate(cell.paragraphs, start=1):
                    table_cell_paragraph_count += 1
                    blocks.append(
                        {
                            "location": (
                                f"table:{table_number}/row:{row_number}/"
                                f"cell:{cell_number}/paragraph:{paragraph_number}"
                            ),
                            "text": paragraph.text,
                        }
                    )
    return blocks, {
        "ordinary_paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "table_cell_paragraph_count": table_cell_paragraph_count,
    }


def extract_proposal_candidates(
    blocks: Iterable[dict[str, str]],
    dictionary_names: Iterable[str],
    structural_names: Iterable[str] = STRUCTURAL_NAMES,
) -> dict[str, dict[str, Any]]:
    dictionary_patterns = {
        name: exact_token_pattern(name) for name in sorted(dictionary_names)
    }
    structural_patterns = {
        name: exact_token_pattern(name) for name in sorted(structural_names)
    }
    candidates: dict[str, dict[str, Any]] = {}

    for block in blocks:
        text = block["text"]
        if not text:
            continue
        names_here = set(RAW_FIELD_TOKEN_RE.findall(text))
        names_here.update(
            name for name, pattern in dictionary_patterns.items() if pattern.search(text)
        )
        names_here.update(
            name for name, pattern in structural_patterns.items() if pattern.search(text)
        )
        for name in sorted(names_here):
            record = candidates.setdefault(
                name, {"locations": [], "contexts": []}
            )
            if block["location"] not in record["locations"]:
                record["locations"].append(block["location"])
            if len(record["contexts"]) < 3:
                excerpt = short_context(text, name)
                if excerpt and excerpt not in record["contexts"]:
                    record["contexts"].append(excerpt)
    return candidates


def read_parquet_metadata(
    path: Path,
    parquet_file_factory: Callable[[Path], Any] = pq.ParquetFile,
) -> dict[str, Any]:
    parquet_file = parquet_file_factory(path)
    metadata = parquet_file.metadata
    schema = parquet_file.schema_arrow
    fields = [
        {"position": position, "name": field.name, "arrow_type": str(field.type)}
        for position, field in enumerate(schema)
    ]
    return {
        "metadata_rows": int(metadata.num_rows),
        "row_group_count": int(metadata.num_row_groups),
        "fields": fields,
    }


def build_table_record(
    path: Path,
    train_dir: Path,
    metadata_reader: Callable[[Path], dict[str, Any]] = read_parquet_metadata,
) -> dict[str, Any]:
    naming = provisional_filename_group(path.name)
    metadata = metadata_reader(path)
    names = [field["name"] for field in metadata["fields"]]
    types = [field["arrow_type"] for field in metadata["fields"]]
    case_type = next(
        (field["arrow_type"] for field in metadata["fields"] if field["name"] == "case_id"),
        None,
    )
    return {
        "file_name": path.name,
        "relative_path": path.relative_to(train_dir).as_posix(),
        "resolved_path": str(path.resolve()),
        "file_size_bytes": int(path.stat().st_size),
        "metadata_rows": metadata["metadata_rows"],
        "row_group_count": metadata["row_group_count"],
        "column_count": len(metadata["fields"]),
        "columns": metadata["fields"],
        "ordered_schema": list(zip(names, types)),
        "case_id_arrow_type": case_type,
        "has_case_id": "case_id" in names,
        "has_target": "target" in names,
        "has_date_decision": "date_decision" in names,
        "has_MONTH": "MONTH" in names,
        "has_WEEK_NUM": "WEEK_NUM" in names,
        "has_num_group1": "num_group1" in names,
        "has_num_group2": "num_group2" in names,
        **naming,
    }


def schema_difference(
    reference: list[tuple[str, str]], candidate: list[tuple[str, str]]
) -> dict[str, Any]:
    reference_types = dict(reference)
    candidate_types = dict(candidate)
    shared = sorted(set(reference_types) & set(candidate_types))
    return {
        "ordered_schema_match": reference == candidate,
        "ordered_field_names_match": [name for name, _ in reference]
        == [name for name, _ in candidate],
        "missing_from_reference_schema": sorted(set(reference_types) - set(candidate_types)),
        "additional_to_reference_schema": sorted(set(candidate_types) - set(reference_types)),
        "type_differences": [
            {
                "field_name": name,
                "reference_type": reference_types[name],
                "candidate_type": candidate_types[name],
            }
            for name in shared
            if reference_types[name] != candidate_types[name]
        ],
    }


def summarize_groups(table_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in table_records:
        grouped[record["provisional_table_group"]].append(record)

    summaries: list[dict[str, Any]] = []
    for group_name in sorted(grouped):
        records = sorted(grouped[group_name], key=lambda item: item["file_name"])
        reference = records[0]["ordered_schema"]
        matches = all(record["ordered_schema"] == reference for record in records)
        differences = []
        if not matches:
            for record in records[1:]:
                difference = schema_difference(reference, record["ordered_schema"])
                if not difference["ordered_schema_match"]:
                    differences.append(
                        {
                            "reference_file": records[0]["file_name"],
                            "compared_file": record["file_name"],
                            **difference,
                        }
                    )
        summary = {
            "provisional_table_group": group_name,
            "physical_files": [record["file_name"] for record in records],
            "physical_file_count": len(records),
            "metadata_physical_row_total": sum(
                record["metadata_rows"] for record in records
            ),
            "column_count_min": min(record["column_count"] for record in records),
            "column_count_max": max(record["column_count"] for record in records),
            "has_case_id_any": any(record["has_case_id"] for record in records),
            "has_case_id_all": all(record["has_case_id"] for record in records),
            "has_num_group1_any": any(record["has_num_group1"] for record in records),
            "has_num_group2_any": any(record["has_num_group2"] for record in records),
            "ordered_schema_match": matches,
            "schema_differences": differences,
        }
        summaries.append(summary)
        for record in records:
            record["group_file_count"] = summary["physical_file_count"]
            record["group_metadata_row_total"] = summary[
                "metadata_physical_row_total"
            ]
            record["group_ordered_schema_match"] = matches
    return summaries


def dictionary_status(
    field_name: str, dictionary: dict[str, str]
) -> tuple[str, str]:
    if field_name in dictionary:
        description = dictionary[field_name]
        return (
            "MATCHED_EMPTY_DESCRIPTION" if description == "" else "MATCHED",
            description,
        )
    if field_name in STRUCTURAL_NAMES:
        return "STRUCTURAL_NOT_IN_DICTIONARY", ""
    return "UNDOCUMENTED", ""


def build_field_inventory(
    table_records: list[dict[str, Any]], dictionary: dict[str, str]
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for table in sorted(table_records, key=lambda item: item["file_name"]):
        for field in table["columns"]:
            match_status, description = dictionary_status(field["name"], dictionary)
            rows.append(
                {
                    "file_name": table["file_name"],
                    "provisional_table_group": table["provisional_table_group"],
                    "column_position_zero_based": field["position"],
                    "field_name": field["name"],
                    "arrow_type": field["arrow_type"],
                    "dictionary_status": match_status,
                    "dictionary_description": description,
                    "technical_role": TECHNICAL_ROLES.get(field["name"], "OTHER"),
                    "substantive_assignment_status": "NOT_ASSIGNED_IN_TASK06",
                }
            )
    return rows


def build_candidate_inventory(
    extracted: dict[str, dict[str, Any]],
    dictionary: dict[str, str],
    field_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    locations: dict[str, dict[str, set[str]]] = defaultdict(
        lambda: {"files": set(), "groups": set(), "types": set()}
    )
    for row in field_rows:
        entry = locations[row["field_name"]]
        entry["files"].add(row["file_name"])
        entry["groups"].add(row["provisional_table_group"])
        entry["types"].add(row["arrow_type"])

    candidate_rows: list[dict[str, Any]] = []
    for name in sorted(extracted):
        observed = name in locations
        files = sorted(locations[name]["files"]) if observed else []
        groups = sorted(locations[name]["groups"]) if observed else []
        types = sorted(locations[name]["types"]) if observed else []
        if not observed:
            location_status = "NOT_OBSERVED_IN_TRAIN"
        elif len(groups) == 1:
            location_status = "OBSERVED_IN_ONE_GROUP"
        else:
            location_status = "OBSERVED_IN_MULTIPLE_GROUPS"
        candidate_rows.append(
            {
                "candidate_name": name,
                "in_dictionary": name in dictionary,
                "dictionary_description": dictionary.get(name, ""),
                "observed_in_train": observed,
                "location_status": location_status,
                "physical_file_count": len(files),
                "provisional_group_count": len(groups),
                "physical_files": files,
                "provisional_groups": groups,
                "arrow_types": types,
                "proposal_reference_locations": list(extracted[name]["locations"]),
                "proposal_contexts": list(extracted[name]["contexts"]),
                "is_structural_reference": name in STRUCTURAL_NAMES,
                "researcher_review_status": "PENDING_RESEARCHER_REVIEW",
            }
        )
    return candidate_rows


def table_csv_row(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "file_name": record["file_name"],
        "relative_path": record["relative_path"],
        "file_size_bytes": record["file_size_bytes"],
        "metadata_rows": record["metadata_rows"],
        "row_group_count": record["row_group_count"],
        "column_count": record["column_count"],
        "provisional_table_group": record["provisional_table_group"],
        "inferred_depth": "" if record["inferred_depth"] is None else record["inferred_depth"],
        "inferred_shard": "" if record["inferred_shard"] is None else record["inferred_shard"],
        "naming_status": record["naming_status"],
        "columns_json": json_text([field["name"] for field in record["columns"]]),
        "arrow_types_json": json_text(
            [field["arrow_type"] for field in record["columns"]]
        ),
        "has_case_id": csv_bool(record["has_case_id"]),
        "case_id_arrow_type": record["case_id_arrow_type"] or "",
        "has_target": csv_bool(record["has_target"]),
        "has_date_decision": csv_bool(record["has_date_decision"]),
        "has_MONTH": csv_bool(record["has_MONTH"]),
        "has_WEEK_NUM": csv_bool(record["has_WEEK_NUM"]),
        "has_num_group1": csv_bool(record["has_num_group1"]),
        "has_num_group2": csv_bool(record["has_num_group2"]),
        "group_file_count": record["group_file_count"],
        "group_metadata_row_total": record["group_metadata_row_total"],
        "group_ordered_schema_match": csv_bool(record["group_ordered_schema_match"]),
    }


def candidate_csv_row(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "candidate_name": record["candidate_name"],
        "in_dictionary": csv_bool(record["in_dictionary"]),
        "dictionary_description": record["dictionary_description"],
        "observed_in_train": csv_bool(record["observed_in_train"]),
        "location_status": record["location_status"],
        "physical_file_count": record["physical_file_count"],
        "provisional_group_count": record["provisional_group_count"],
        "physical_files_json": json_text(record["physical_files"]),
        "provisional_groups_json": json_text(record["provisional_groups"]),
        "arrow_types_json": json_text(record["arrow_types"]),
        "proposal_reference_locations_json": json_text(
            record["proposal_reference_locations"]
        ),
        "proposal_contexts_json": json_text(record["proposal_contexts"]),
        "is_structural_reference": csv_bool(record["is_structural_reference"]),
        "researcher_review_status": record["researcher_review_status"],
    }


def write_csv_rows(handle: TextIO, headers: Iterable[str], rows: Iterable[dict[str, Any]]) -> None:
    writer = csv.DictWriter(handle, fieldnames=list(headers), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)


def write_csv_file(path: Path, headers: Iterable[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        write_csv_rows(handle, headers, rows)


def validate_output_targets(output_dir: Path) -> dict[str, Path]:
    targets = {name: output_dir / name for name in REPORT_FILENAMES}
    existing = [str(path) for path in targets.values() if path.exists()]
    if existing:
        raise FileExistsError(f"Refusing to overwrite existing report targets: {existing}")
    return targets


def collect_input_state(paths: Iterable[Path]) -> list[dict[str, Any]]:
    state = []
    for path in sorted(paths, key=lambda item: str(item.resolve())):
        stat = path.stat()
        state.append(
            {
                "resolved_path": str(path.resolve()),
                "size_bytes": int(stat.st_size),
                "mtime_ns": int(stat.st_mtime_ns),
            }
        )
    return state


def compare_input_states(
    before: list[dict[str, Any]], after: list[dict[str, Any]]
) -> dict[str, Any]:
    before_by_path = {row["resolved_path"]: row for row in before}
    after_by_path = {row["resolved_path"]: row for row in after}
    all_paths = sorted(set(before_by_path) | set(after_by_path))
    comparisons = []
    for path in all_paths:
        before_row = before_by_path.get(path)
        after_row = after_by_path.get(path)
        unchanged = before_row == after_row
        comparisons.append(
            {
                "resolved_path": path,
                "unchanged": unchanged,
                "before": before_row,
                "after": after_row,
            }
        )
    return {
        "method": "file size and mtime_ns comparison; not a cryptographic guarantee",
        "all_unchanged": all(row["unchanged"] for row in comparisons),
        "comparisons": comparisons,
    }


def build_internal_checks(
    table_records: list[dict[str, Any]],
    field_rows: list[dict[str, Any]],
    candidate_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    expected_field_occurrences = sum(record["column_count"] for record in table_records)
    candidate_names = [row["candidate_name"] for row in candidate_rows]
    status_counts = Counter(row["location_status"] for row in candidate_rows)
    observed_candidates_resolved = all(
        row["physical_file_count"] > 0 and row["provisional_group_count"] > 0
        for row in candidate_rows
        if row["observed_in_train"]
    )
    checks = {
        "field_occurrence_count_reconciles": len(field_rows) == expected_field_occurrences,
        "expected_field_occurrence_count": expected_field_occurrences,
        "actual_field_occurrence_count": len(field_rows),
        "proposal_candidate_names_unique": len(candidate_names) == len(set(candidate_names)),
        "candidate_status_counts_reconcile": sum(status_counts.values())
        == len(candidate_rows),
        "observed_candidates_have_schema_locations": observed_candidates_resolved,
    }
    checks["all_internal_checks_pass"] = all(
        value
        for key, value in checks.items()
        if key.endswith("_reconciles")
        or key.endswith("_unique")
        or key.endswith("_locations")
    )
    return checks


def build_summary(
    args: argparse.Namespace,
    targets: dict[str, Path],
    train_files: list[Path],
    before_state: list[dict[str, Any]],
    after_state: list[dict[str, Any]],
    dictionary_rows: list[dict[str, str]],
    dictionary: dict[str, str],
    dictionary_findings: dict[str, Any],
    table_records: list[dict[str, Any]],
    group_summaries: list[dict[str, Any]],
    field_rows: list[dict[str, Any]],
    candidate_rows: list[dict[str, Any]],
    proposal_coverage: dict[str, int],
) -> dict[str, Any]:
    distinct_fields = sorted({row["field_name"] for row in field_rows})
    dictionary_names = set(dictionary)
    observed_fields = set(distinct_fields)
    dictionary_observed = sorted(dictionary_names & observed_fields)
    dictionary_not_observed = sorted(dictionary_names - observed_fields)
    undocumented_nonstructural = sorted(
        {
            row["field_name"]
            for row in field_rows
            if row["dictionary_status"] == "UNDOCUMENTED"
        }
    )
    structural_not_documented = sorted(
        {
            row["field_name"]
            for row in field_rows
            if row["dictionary_status"] == "STRUCTURAL_NOT_IN_DICTIONARY"
        }
    )
    groups_by_field: dict[str, set[str]] = defaultdict(set)
    for row in field_rows:
        groups_by_field[row["field_name"]].add(row["provisional_table_group"])
    fields_multiple_groups = sorted(
        name for name, groups in groups_by_field.items() if len(groups) > 1
    )
    matching_status_counts = Counter(row["dictionary_status"] for row in field_rows)
    documented_distinct = sorted(observed_fields & dictionary_names)
    candidate_status_counts = Counter(row["location_status"] for row in candidate_rows)
    candidate_absent = sorted(
        row["candidate_name"]
        for row in candidate_rows
        if row["location_status"] == "NOT_OBSERVED_IN_TRAIN"
    )
    candidate_multiple_groups = sorted(
        row["candidate_name"]
        for row in candidate_rows
        if row["location_status"] == "OBSERVED_IN_MULTIPLE_GROUPS"
    )
    naming_warnings = sorted(
        record["naming_warning"]
        for record in table_records
        if record["naming_warning"] is not None
    )
    schema_warning_groups = sorted(
        group["provisional_table_group"]
        for group in group_summaries
        if not group["ordered_schema_match"]
    )
    input_comparison = compare_input_states(before_state, after_state)
    internal_checks = build_internal_checks(table_records, field_rows, candidate_rows)

    issues = []
    if len(train_files) != EXPECTED_TRAIN_FILE_COUNT:
        issues.append(
            f"Observed {len(train_files)} TRAIN Parquet files; expected {EXPECTED_TRAIN_FILE_COUNT}."
        )
    if naming_warnings:
        issues.append(f"Unexpected filename patterns: {len(naming_warnings)}")
    if schema_warning_groups:
        issues.append(
            f"Provisional groups with ordered-schema differences: {schema_warning_groups}"
        )
    if not input_comparison["all_unchanged"]:
        issues.append("At least one input size or modification time changed during execution.")
    if not internal_checks["all_internal_checks_pass"]:
        issues.append("At least one internal report-consistency check failed.")

    status = "COMPLETED_WITH_WARNINGS" if issues else "COMPLETED"
    return {
        "task_identifier": TASK_ID,
        "script_version": SCRIPT_VERSION,
        "utc_execution_timestamp": datetime.now(timezone.utc).isoformat().replace(
            "+00:00", "Z"
        ),
        "runtime_versions": {
            "python": platform.python_version(),
            "pyarrow": pyarrow.__version__,
            "python_docx": docx.__version__,
        },
        "resolved_paths": {
            "train_directory": str(args.train_dir.resolve()),
            "dictionary": str(args.dictionary.resolve()),
            "proposal": str(args.proposal.resolve()),
            "output_directory": str(args.output_dir.resolve()),
            "outputs": {name: str(path.resolve()) for name, path in targets.items()},
        },
        "source_inventory": {
            "observed_train_parquet_file_count": len(train_files),
            "expected_train_parquet_file_count": EXPECTED_TRAIN_FILE_COUNT,
            "matches_expected_file_count": len(train_files)
            == EXPECTED_TRAIN_FILE_COUNT,
            "provisional_group_count": len(group_summaries),
            "physical_metadata_row_total": sum(
                record["metadata_rows"] for record in table_records
            ),
            "physical_metadata_row_total_interpretation": (
                "Sum of physical-file metadata row counts; not a distinct-application count."
            ),
            "total_field_occurrence_count": len(field_rows),
            "distinct_field_name_count": len(distinct_fields),
        },
        "dictionary": {
            "row_count": len(dictionary_rows),
            "unique_variable_count": len(dictionary),
            "validation_findings": dictionary_findings,
        },
        "dictionary_schema_matching": {
            "field_occurrence_counts_by_status": dict(sorted(matching_status_counts.items())),
            "documented_distinct_observed_field_count": len(documented_distinct),
            "undocumented_nonstructural_distinct_field_count": len(
                undocumented_nonstructural
            ),
            "structural_not_in_dictionary_distinct_field_count": len(
                structural_not_documented
            ),
            "distinct_observed_field_names": distinct_fields,
            "dictionary_variables_observed_in_train_schemas": dictionary_observed,
            "dictionary_variables_not_observed_in_train_schemas": dictionary_not_observed,
            "undocumented_nonstructural_train_fields": undocumented_nonstructural,
            "structural_train_fields_not_in_dictionary": structural_not_documented,
            "distinct_field_names_in_multiple_provisional_groups": fields_multiple_groups,
            "not_observed_interpretation": (
                "Not observed in TRAIN schemas does not mean a dictionary entry is invalid "
                "or unavailable everywhere."
            ),
        },
        "proposal_extraction": {
            **proposal_coverage,
            "method": (
                "python-docx ordinary paragraphs and table-cell paragraphs with exact "
                "ASCII letter/digit/underscore token boundaries"
            ),
            "candidate_count": len(candidate_rows),
            "candidate_counts_by_location_status": dict(
                sorted(candidate_status_counts.items())
            ),
            "candidate_names_not_observed_in_train": candidate_absent,
            "candidate_names_observed_in_multiple_groups": candidate_multiple_groups,
            "coverage_limitations": [
                "No visual rendering or complete visual review was performed.",
                "Only ordinary document paragraphs and table-cell paragraph text were inspected.",
                "Images, text boxes, headers, footers, and other non-paragraph content were not extracted.",
                "A proposal reference is not treated as approval for predictive use.",
            ],
        },
        "provisional_group_findings": group_summaries,
        "naming_warnings": naming_warnings,
        "input_before_after_comparison": input_comparison,
        "internal_reconciliation_checks": internal_checks,
        "explicitly_unperformed_checks": list(UNPERFORMED_CHECKS),
        "report_status": {
            "execution_status": status,
            "internal_report_consistency": (
                "PASS" if internal_checks["all_internal_checks_pass"] else "FAIL"
            ),
            "observed_warnings": issues,
            "unassessed_substantive_questions": "NOT_ASSESSED",
            "note": (
                "Successful inventory execution does not mean the data are model-ready "
                "or that any field is eligible for modeling."
            ),
        },
    }


def main() -> int:
    args = parse_args()
    for label, path in (
        ("TRAIN directory", args.train_dir),
        ("dictionary", args.dictionary),
        ("proposal", args.proposal),
    ):
        if not path.exists():
            raise FileNotFoundError(f"Required {label} does not exist: {path}")
    if not args.train_dir.is_dir():
        raise NotADirectoryError(f"TRAIN path is not a directory: {args.train_dir}")
    if not args.dictionary.is_file() or not args.proposal.is_file():
        raise FileNotFoundError("Dictionary and proposal paths must be files")

    targets = validate_output_targets(args.output_dir)
    train_files = sorted(
        (
            path
            for path in args.train_dir.iterdir()
            if path.is_file() and path.suffix == ".parquet"
        ),
        key=lambda path: path.name,
    )
    input_paths = [*train_files, args.dictionary, args.proposal]
    before_state = collect_input_state(input_paths)

    dictionary_rows, dictionary, dictionary_findings = read_dictionary(
        args.dictionary
    )
    table_records = [
        build_table_record(path, args.train_dir) for path in train_files
    ]
    group_summaries = summarize_groups(table_records)
    field_rows = build_field_inventory(table_records, dictionary)

    proposal_document = docx.Document(args.proposal)
    blocks, proposal_coverage = proposal_text_blocks(proposal_document)
    extracted = extract_proposal_candidates(blocks, dictionary, STRUCTURAL_NAMES)
    candidate_rows = build_candidate_inventory(extracted, dictionary, field_rows)
    after_state = collect_input_state(input_paths)

    summary = build_summary(
        args,
        targets,
        train_files,
        before_state,
        after_state,
        dictionary_rows,
        dictionary,
        dictionary_findings,
        table_records,
        group_summaries,
        field_rows,
        candidate_rows,
        proposal_coverage,
    )
    if not summary["internal_reconciliation_checks"]["all_internal_checks_pass"]:
        raise RuntimeError("Internal report-consistency checks failed; no reports written")

    validate_output_targets(args.output_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv_file(
        targets["table_inventory.csv"],
        TABLE_HEADERS,
        [table_csv_row(record) for record in table_records],
    )
    write_csv_file(targets["field_inventory.csv"], FIELD_HEADERS, field_rows)
    write_csv_file(
        targets["proposal_candidates.csv"],
        CANDIDATE_HEADERS,
        [candidate_csv_row(record) for record in candidate_rows],
    )
    targets["inventory_summary.json"].write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False) + "\n",
        encoding="utf-8",
    )

    print(
        json.dumps(
            {
                "execution_status": summary["report_status"]["execution_status"],
                "train_parquet_file_count": len(table_records),
                "provisional_group_count": len(group_summaries),
                "field_occurrence_count": len(field_rows),
                "proposal_candidate_count": len(candidate_rows),
                "output_paths": [str(targets[name]) for name in REPORT_FILENAMES],
            },
            indent=2,
            ensure_ascii=False,
            allow_nan=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
